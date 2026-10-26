"""Draft pptx of one figure: places its atoms and text boxes by experiments/<fig>/layout.py (RESULTS_PLAN.md §2).

Usage (login node): PYTHONPATH=experiments/common/.pylib python experiments/common/build_figure.py <fig>
Writes figures/<fig>/<fig>_draft.pptx only (overwritten every run). The user copies it to figures/<fig>/<fig>.pptx and
edits that copy, which no script ever writes.

experiments/<fig>/layout.py defines:
  HEIGHT_MM  slide height (the width is the 180 mm page)
  ATOMS      [(atom, x_mm, y_mm, scale)]: top-left corner; the size is the manifest size times scale, and any scale
             other than 1 is reported (redraw the atom at its placed size). x >= 180 parks a spare atom off the page.
  TEXTS      [(text, x_mm, y_mm, w_mm, h_mm, kind)]: kind "letter" (panel letter, 8 pt bold) or "label" (7 pt) or
             "small" (6 pt); every text box stays editable.
- Atoms are the 600 dpi PNGs (figures/<fig>/atoms/png); the vector PDFs stay beside them.
- python-pptx lives in experiments/common/.pylib (pip --target), so spindle_env stays untouched.
"""

import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import style as ps  # noqa: E402

FONT = "Arial"  # metric twin of Liberation Sans (the atoms' font)
INK = RGBColor(0x1A, 0x1A, 0x1A)
KINDS = {"letter": (8, True), "label": (ps.LABEL_PT, False), "small": (ps.TICK_PT, False)}


def load_layout(fig):
    path = ps.ROOT / "experiments" / fig / "layout.py"
    spec = importlib.util.spec_from_file_location(f"layout_{fig}", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def text_box(slide, text, x, y, w, h, kind):
    pt, bold = KINDS[kind]
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.name, r.font.bold, r.font.color.rgb = Pt(pt), FONT, bold, INK
    return tb


def build(fig):
    lay = load_layout(fig)
    sizes = ps.manifest()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(ps.PAGE_W), Mm(lay.HEIGHT_MM)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    placed = set()
    for name, x, y, scale in lay.ATOMS:
        w, h = sizes[fig, name]
        png = ps.atom_paths(fig, name)[1]
        slide.shapes.add_picture(str(png), Mm(x), Mm(y), Mm(w * scale), Mm(h * scale))
        placed.add(name)
        if scale != 1:
            print(f"WARNING scaled: {name} x {scale} (placed {w * scale:.1f} x {h * scale:.1f} mm, drawn {w:.1f} x "
                  f"{h:.1f}); redraw it at its placed size")
        if x < ps.PAGE_W and (x + w * scale > ps.PAGE_W + 0.01 or y + h * scale > lay.HEIGHT_MM + 0.01):
            print(f"WARNING: {name} runs past the page")
    for text, x, y, w, h, kind in lay.TEXTS:
        text_box(slide, text, x, y, w, h, kind)
    missing = {a for f, a in sizes if f == fig} - placed
    if missing:
        print("atoms not placed:", ", ".join(sorted(missing)))
    out = ps.FIGURES / fig / f"{fig}_draft.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"wrote {out.relative_to(ps.ROOT)} (180 x {lay.HEIGHT_MM:g} mm)")
    return out


if __name__ == "__main__":
    build(sys.argv[1])
