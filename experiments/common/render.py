"""Render a figure's edited pptx: figures/<fig>/<fig>.pptx -> figures/<fig>/<fig>.pdf + <fig>.png.

Usage (login node): PYTHONPATH=experiments/common/.pylib python experiments/common/render.py <fig> [--draft]
--draft renders <fig>_draft.pptx instead (to <fig>_draft.pdf / .png), for a look before the user's copy exists.

Resolution rule (user, 26 Oct): nothing may break when zoomed in, within the submission size budget. So the PDF is
vector: every picture that is an atom (its bytes equal an atom PNG of this figure, else its name is an atom's name)
is taken out of a copy of the deck, LibreOffice (headless, a throwaway profile) exports what is left (text boxes,
lines, any other picture, lossless at full resolution), and each atom's vector PDF (transparent) is stamped onto that
page at the picture's place and size. A cropped or rotated atom picture stays the embedded 1200 dpi PNG (reported).
pdftoppm rasterises the PDF to a 1200 dpi PNG. The PDF page must be 180 mm wide (the slide size); a file over
BUDGET_MB is reported.
"""

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader, PdfWriter, Transformation

ROOT = Path(__file__).resolve().parents[2]
FIGURES = ROOT / "figures"
DPI = 1200
BUDGET_MB = 10.0  # per figure file (journal figure-file limit)
EMU_PER_PT = 12700
PDF_FILTER = "pdf:impress_pdf_Export:" + json.dumps({
    "ReduceImageResolution": {"type": "boolean", "value": "false"},
    "UseLosslessCompression": {"type": "boolean", "value": "true"},
    "ExportNotesPages": {"type": "boolean", "value": "false"},
})


def soffice_pdf(pptx, out_pdf):
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp / 'profile'}", "--headless", "--convert-to",
                        PDF_FILTER, "--outdir", str(tmp), str(pptx)], check=True, capture_output=True)
        shutil.move(str(tmp / f"{pptx.stem}.pdf"), str(out_pdf))


def atom_index(fig):
    """{sha1 of the atom PNG: name} and {name: vector PDF} of the figure's atoms."""
    d = FIGURES / fig / "atoms"
    by_hash = {hashlib.sha1(p.read_bytes()).hexdigest(): p.stem for p in (d / "png").glob("*.png")}
    pdfs = {p.stem: p for p in (d / "pdf").glob("*.pdf")}
    return by_hash, pdfs


def split_deck(fig, pptx, tmp):
    """A copy of the deck without its atom pictures, and [(atom, left, top, width, height) in pt] of those removed."""
    by_hash, pdfs = atom_index(fig)
    prs = Presentation(str(pptx))
    slide = prs.slides[0]
    placed = []
    for shape in list(slide.shapes):
        if shape.shape_type != 13:  # MSO_SHAPE_TYPE.PICTURE
            continue
        name = by_hash.get(hashlib.sha1(shape.image.blob).hexdigest())
        if name is None and shape.name in pdfs:
            name = shape.name
            print(f"note: {name}: the picture is an older drawing of the atom; its current vector PDF is used")
        if name is None or name not in pdfs:
            continue
        crop = any(abs(getattr(shape, f"crop_{s}")) > 1e-6 for s in ("left", "right", "top", "bottom"))
        if crop or shape.rotation:
            print(f"WARNING: {name} is {'cropped' if crop else 'rotated'}; it stays a 1200 dpi picture")
            continue
        placed.append((name, *(v / EMU_PER_PT for v in (shape.left, shape.top, shape.width, shape.height))))
        shape._element.getparent().remove(shape._element)
    stripped = tmp / pptx.name
    prs.save(stripped)
    return stripped, placed, pdfs, prs.slide_height / EMU_PER_PT


def to_pdf(fig, pptx, out_pdf):
    with tempfile.TemporaryDirectory() as tmp:
        stripped, placed, pdfs, page_h = split_deck(fig, pptx, Path(tmp))
        base = Path(tmp) / "base.pdf"
        soffice_pdf(stripped, base)
        page = PdfReader(base).pages[0]
        for name, x, y, w, h in placed:
            atom = PdfReader(pdfs[name]).pages[0]
            aw, ah = float(atom.mediabox.width), float(atom.mediabox.height)
            sx, sy = w / aw, h / ah
            if abs(sx - 1) > 0.01 or abs(sy - 1) > 0.01:
                print(f"WARNING scaled: {name} x {sx:.3f} / {sy:.3f}; redraw it at its placed size")
            page.merge_transformed_page(atom, Transformation().scale(sx, sy).translate(x, page_h - y - h))
        out = PdfWriter()
        out.add_page(page)
        out.compress_identical_objects()
        with open(out_pdf, "wb") as fh:
            out.write(fh)
    w_mm, h_mm = (float(v) / 72 * 25.4 for v in (page.mediabox.width, page.mediabox.height))
    if abs(w_mm - 180) > 0.5:
        raise RuntimeError(f"{out_pdf.name}: page is {w_mm:.1f} mm wide, not 180 mm")
    return w_mm, h_mm, len(placed)


def to_png(pdf, out_png):
    subprocess.run(["pdftoppm", "-png", "-r", str(DPI), "-singlefile", str(pdf), str(out_png.with_suffix(""))],
                   check=True)


def check_size(*paths):
    for p in paths:
        mb = p.stat().st_size / 2**20
        print(f"  {p.relative_to(ROOT)}: {mb:.1f} MB" + (f"  OVER the {BUDGET_MB:g} MB budget" if mb > BUDGET_MB else ""))


def render(fig, draft=False):
    stem = f"{fig}_draft" if draft else fig
    pptx = FIGURES / fig / f"{stem}.pptx"
    if not pptx.exists():
        raise SystemExit(f"{pptx.relative_to(ROOT)} does not exist" +
                         ("" if draft else f" (copy {fig}_draft.pptx to it and edit the copy)"))
    pdf, png = pptx.with_suffix(".pdf"), pptx.with_suffix(".png")
    w, h, n = to_pdf(fig, pptx, pdf)
    to_png(pdf, png)
    print(f"wrote {pdf.name} ({n} vector atoms) and {png.name} ({w:.0f} x {h:.0f} mm, {DPI} dpi)")
    check_size(pptx, pdf, png)


if __name__ == "__main__":
    render(sys.argv[1], draft="--draft" in sys.argv)
